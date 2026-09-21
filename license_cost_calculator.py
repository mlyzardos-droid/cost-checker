#!/usr/bin/env python3
"""Monthly AI licence cost calculator for the AD Based License Report.

Run with: python3 license_cost_calculator.py
The application opens the source Excel report and asks where to save the export.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
import json
from pathlib import Path
import random
import shutil
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit(
        "Missing package 'openpyxl'. Install it once with: python3 -m pip install openpyxl"
    )


# Monthly prices in EUR per assigned licence/user. ChatGPT and GitHub package
# prices come from the supplied images.
PRODUCT_PRICES = {
    "ChatGPT Enterprise Basic": ("ChatGPT Enterprise Basic", 0.00),
    "ChatGPT Enterprise Credit Package L": ("ChatGPT Credit Package L", 100.00),
    "ChatGPT Enterprise Credit Package M": ("ChatGPT Credit Package M", 30.00),
    "ChatGPT Enterprise Credit Package XXL": ("ChatGPT Credit Package XXL", 1000.00),
    "Gemini Code Assist": ("Gemini", 20.00),
    "M365 Copilot License": ("MS Copilot", 5.70),
    "Claude Code": ("Claude Code", 200.00),
    "Cursor": ("Cursor", 40.00),
    "GitHub Copilot": ("GitHub Copilot", 20.00),
    "GitHub Copilot Upgrade Package 1 for advanced use": ("GitHub Copilot Upgrade - Advanced", 100.00),
    "GitHub Copilot Upgrade Package 2 for intensive use": ("GitHub Copilot Upgrade - Intensive", 250.00),
}
NON_AI_PRODUCT_PRICES = {
    "Bitwarden": ("Bitwarden", 0.60),
    "Miro Enterprise License": ("Miro Enterprise", 9.00),
    "Global Secure Access for LMD": ("Global Secure Access for LMD", 5.10),
}
DEFAULT_PRODUCT_PRICES = dict(PRODUCT_PRICES)
LICENSE_ACTIVE = {name: True for name in PRODUCT_PRICES}
NON_AI_LICENSE_ACTIVE = {name: True for name in NON_AI_PRODUCT_PRICES}

def export_directory() -> Path:
    """Find the user's Desktop/AI Cost folder, including common OneDrive layouts."""
    home = Path.home()
    candidates = [
        home / "Desktop" / "AI Cost",
        home / "OneDrive" / "Desktop" / "AI Cost",
    ]
    candidates.extend(home.glob("OneDrive -*/*/AI Cost"))
    for candidate in candidates:
        if candidate.is_dir():
            return candidate
    candidates[0].mkdir(parents=True, exist_ok=True)
    return candidates[0]


def history_file() -> Path:
    folder = export_directory() / "history"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / "AI_Cost_History.xlsx"


def members_file() -> Path:
    folder = export_directory() / "history"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / "costcenter_members.json"


def license_settings_file() -> Path:
    folder = export_directory() / "history"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / "license_settings.json"


def load_license_settings():
    """Load the current licence catalogue without changing old exports/history."""
    path = license_settings_file()
    if not path.exists():
        return
    try:
        with path.open("r", encoding="utf-8") as file:
            settings = json.load(file)
        sections = settings if any(key in settings for key in ("ai", "non_ai")) else {"ai": settings}
        for name, values in sections.get("ai", {}).items():
            if not isinstance(values, dict):
                continue
            price = number(values.get("price"))
            if price is None:
                continue
            product = clean(values.get("product")) or name
            PRODUCT_PRICES[name] = (product, price)
            LICENSE_ACTIVE[name] = bool(values.get("active", True))
        for name, values in sections.get("non_ai", {}).items():
            if not isinstance(values, dict):
                continue
            price = number(values.get("price"))
            if price is None:
                continue
            product = clean(values.get("product")) or name
            NON_AI_PRODUCT_PRICES[name] = (product, price)
            NON_AI_LICENSE_ACTIVE[name] = bool(values.get("active", True))
    except (OSError, json.JSONDecodeError):
        return


def save_license_settings():
    settings = {
        "ai": {
            name: {"product": product, "price": price, "active": LICENSE_ACTIVE.get(name, True)}
            for name, (product, price) in PRODUCT_PRICES.items()
        },
        "non_ai": {
            name: {"product": product, "price": price, "active": NON_AI_LICENSE_ACTIVE.get(name, True)}
            for name, (product, price) in NON_AI_PRODUCT_PRICES.items()
        },
    }
    with license_settings_file().open("w", encoding="utf-8") as file:
        json.dump(settings, file, ensure_ascii=False, indent=2)


def active_license_names():
    return {name for name in PRODUCT_PRICES if LICENSE_ACTIVE.get(name, True)}


def active_non_ai_license_names():
    return {name for name in NON_AI_PRODUCT_PRICES if NON_AI_LICENSE_ACTIVE.get(name, True)}


def load_saved_members():
    path = members_file()
    if not path.exists():
        return {}, "Not configured"
    try:
        with path.open("r", encoding="utf-8") as file:
            members = json.load(file)
        updated = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
        return members, updated
    except (OSError, json.JSONDecodeError):
        return {}, "Not configured"


def save_members(members):
    path = members_file()
    with path.open("w", encoding="utf-8") as file:
        json.dump(members, file, ensure_ascii=False, indent=2)
    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M")


def clean(value) -> str:
    return str(value).strip() if value is not None else ""


def number(value):
    if isinstance(value, (int, float)):
        return value
    text = clean(value).replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def read_report(source: Path):
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    # The supplied report uses columns B, E, F, G and H for the useful fields.
    services = []
    current = None
    for row in rows:
        row = list(row) + [None] * 8
        service_name = clean(row[1])
        total = number(row[5])
        if service_name and total is not None:
            current = {
                "service": service_name,
                "material": clean(row[4]),
                "reported_total": int(total) if float(total).is_integer() else total,
                "users": [],
            }
            services.append(current)
        elif current:
            account = clean(row[7])
            username = clean(row[6])
            if account or username:
                current["users"].append(
                    {"account": account, "username": username}
                )

    return services


def read_costcenter_report(source: Path):
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    members = {}
    for row in rows[5:]:
        row = list(row) + [None] * 11
        account = clean(row[6]).casefold()
        if account:
            members[account] = {
                "costcenter": clean(row[1]),
                "name": clean(row[4]),
                "email": clean(row[5]),
                "account": clean(row[6]),
                "manager": clean(row[8]),
                "manager_account": clean(row[9]),
            }
    return members


def calculate(services, members=None, include_non_ai=False):
    members = members or {}
    catalog = dict(PRODUCT_PRICES)
    active_names = active_license_names()
    if include_non_ai:
        catalog.update(NON_AI_PRODUCT_PRICES)
        active_names |= active_non_ai_license_names()
    detail = []
    unique_users = set()
    totals = defaultdict(lambda: {"users": set(), "licenses": 0, "cost": 0.0, "priced": True})
    user_costs = {}
    user_licenses = []

    for item in services:
        if item["service"] not in active_names:
            continue
        product, price = catalog[item["service"]]
        accounts = {
            u["account"] or u["username"]
            for u in item["users"]
            if u["account"] or u["username"]
        }
        licenses = item["reported_total"]
        cost = float(licenses) * price if price is not None else None
        unique_users.update(accounts)
        totals[product]["users"].update(accounts)
        totals[product]["licenses"] += licenses
        if cost is None:
            totals[product]["priced"] = False
        else:
            totals[product]["cost"] += cost
        detail.append(
            {
                "product": product,
                "service": item["service"],
                "material": item["material"],
                "licenses": licenses,
                "price": price,
                "cost": cost,
                "users": len(accounts),
            }
        )
        for user in item["users"]:
            account = user["account"] or user["username"]
            if not account:
                continue
            key = account.casefold()
            record = user_costs.setdefault(key, {
                "account": account, "name": user["username"], "email": "",
                "costcenter": "", "manager": "", "licenses": 0, "cost": 0.0,
                "products": set(), "chargeable_licenses": 0,
            })
            member = members.get(key, {})
            record["name"] = member.get("name") or record["name"]
            record["email"] = member.get("email", "")
            record["costcenter"] = member.get("costcenter", "")
            record["manager"] = member.get("manager", "")
            record["licenses"] += 1
            record["products"].add(product)
            if price is not None and price > 0:
                record["chargeable_licenses"] += 1
            if price is not None:
                record["cost"] += price
            user_licenses.append({
                "account": account,
                "name": record["name"],
                "email": record["email"],
                "costcenter": record["costcenter"],
                "manager": record["manager"],
                "product": product,
                "service": item["service"],
                "price": price,
                "cost": price,
            })

    summary = []
    for product, values in totals.items():
        summary.append(
            {
                "product": product,
                "licenses": values["licenses"],
                "unique_users": len(values["users"]),
                "cost": values["cost"] if values["priced"] else None,
                "priced": values["priced"],
            }
        )
    summary.sort(key=lambda x: x["product"])
    without_license = [
        member for key, member in members.items()
        if key not in user_costs
    ]
    without_license.sort(key=lambda x: x.get("account", "").casefold())
    return (
        summary,
        detail,
        unique_users,
        sorted(user_costs.values(), key=lambda x: x["account"].casefold()),
        without_license,
        user_licenses,
    )


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="E20074")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        width = min(max(len(clean(c.value)) for c in column_cells) + 2, 40)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def export_report(destination: Path, source: Path, summary, detail, unique_users, user_costs, without_license, user_licenses, missing_emails=None, full_license_report=False):
    wb = Workbook()
    overview = wb.active
    overview.title = "Summary"
    overview.append(["Metric", "Value"])
    total_cost = sum(row["cost"] or 0 for row in summary)
    overview.append(["Source report", source.name])
    overview.append(["Created", datetime.now().strftime("%Y-%m-%d %H:%M")])
    overview.append(["Unique licensed users" if full_license_report else "Unique AI users", len(unique_users)])
    overview.append(["Total monthly cost (EUR)", total_cost])
    overview.append(["12-month projection (EUR)", total_cost * 12])
    overview.append([
        "Average cost per unique AI user (EUR)",
        total_cost / len(unique_users) if unique_users else "n.a.",
    ])
    overview.append([])
    overview.append(["Product", "Assigned licences", "Unique users", "Monthly cost (EUR)", "12-month projection (EUR)"])
    for row in summary:
        overview.append([row["product"], row["licenses"], row["unique_users"], row["cost"], row["cost"] * 12])
    style_sheet(overview)
    for row_number in (5, 6, 7):
        overview.cell(row_number, 2).number_format = '€#,##0.00'
    for cell in overview[8]:
        cell.fill = PatternFill("solid", fgColor="E20074")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in overview.iter_rows(min_row=9, min_col=4, max_col=5):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '€#,##0.00'

    detail_sheet = wb.create_sheet("All Licenses" if full_license_report else "Details")
    detail_headers = ["Product", "Report service", "Material number", "Assigned licences", "Price/user/month (EUR)", "Monthly cost (EUR)", "Users found"]
    if full_license_report:
        detail_headers.insert(1, "License type")
    detail_sheet.append(detail_headers)
    for row in detail:
        values = [row["product"], row["service"], row["material"], row["licenses"], row["price"], row["cost"], row["users"]]
        if full_license_report:
            values.insert(1, "AI" if row["service"] in PRODUCT_PRICES else "Non-AI")
        detail_sheet.append(values)
    style_sheet(detail_sheet)
    price_column = 6 if full_license_report else 5
    cost_column = 7 if full_license_report else 6
    for row in detail_sheet.iter_rows(min_row=2, min_col=price_column, max_col=cost_column):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '€#,##0.00'

    user_cost_sheet = wb.create_sheet("User Costs")
    user_cost_sheet.append([
        "Account", "Name", "Email", "Cost center", "Cost center manager",
        "All licenses" if full_license_report else "AI licences", "Monthly cost (EUR)",
    ])
    for user in user_costs:
        user_cost_sheet.append([
            user["account"], user["name"], user["email"], user["costcenter"],
            user["manager"], user["licenses"], user["cost"],
        ])
    style_sheet(user_cost_sheet)
    for row in user_cost_sheet.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = '€#,##0.00'

    user_cost_sheet.delete_rows(1, user_cost_sheet.max_row)
    user_cost_sheet.append([
        "Account", "Name", "Email", "Cost center", "Cost center manager",
        "Licenses assigned", "AI licences", "Monthly cost (EUR)",
        "Chargeable licences", "12-month projection (EUR)",
    ])
    for user in user_costs:
        user_cost_sheet.append([
            user["account"], user["name"], user["email"], user["costcenter"],
            user["manager"], "; ".join(sorted(user["products"])), user["licenses"],
            user["cost"], user["chargeable_licenses"], user["cost"] * 12,
        ])
    style_sheet(user_cost_sheet)
    for row in user_cost_sheet.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            if cell.column in (8, 10) and isinstance(cell.value, (int, float)):
                cell.number_format = '€#,##0.00'

    without_sheet = wb.create_sheet("Users Without Any License" if full_license_report else "Users Without AI License")
    without_sheet.append([
        "Account", "Name", "Email", "Cost center", "Cost center manager",
    ])
    for member in without_license:
        without_sheet.append([
            member.get("account", ""), member.get("name", ""),
            member.get("email", ""), member.get("costcenter", ""),
            member.get("manager", ""),
        ])
    style_sheet(without_sheet)

    if missing_emails:
        missing_sheet = wb.create_sheet("Requested Emails Not Found")
        missing_sheet.append(["Email"])
        for email in sorted(missing_emails):
            missing_sheet.append([email])
        style_sheet(missing_sheet)

    wb.save(destination)


def append_history(
    source: Path, summary, unique_users, without_license_count, member_updated,
    history_sheet_name="History", product_sheet_name="Product History",
    full_license_report=False,
):
    history_path = history_file()
    if history_path.exists():
        wb = load_workbook(history_path)
        sheet = wb[history_sheet_name] if history_sheet_name in wb.sheetnames else wb.create_sheet(history_sheet_name)
    else:
        wb = Workbook()
        sheet = wb.active
        sheet.title = history_sheet_name
    if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
        sheet.append([
            "Date/time", "Source report",
            "Unique licensed users" if full_license_report else "Unique AI users",
            "Monthly cost (EUR)", "12-month projection (EUR)", "Average cost/user (EUR)",
            "Costcenter data updated",
            "Users without any license" if full_license_report else "Users without AI license",
        ])
    sheet.cell(1, 7).value = "Costcenter data updated"
    sheet.cell(1, 8).value = "Users without any license" if full_license_report else "Users without AI license"
    total_cost = sum(row["cost"] or 0 for row in summary)
    sheet.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        source.name,
        len(unique_users),
        total_cost,
        total_cost * 12,
        total_cost / len(unique_users) if unique_users else "n.a.",
        member_updated,
        without_license_count,
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E20074")
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '€#,##0.00'
    for column_cells in sheet.columns:
        width = min(max(len(clean(c.value)) for c in column_cells) + 2, 45)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    sheet.freeze_panes = "A2"
    wb.save(history_path)


    product_sheet = wb[product_sheet_name] if product_sheet_name in wb.sheetnames else wb.create_sheet(product_sheet_name)
    if product_sheet.max_row == 1 and product_sheet.cell(1, 1).value is None:
        product_sheet.delete_rows(1, 1)
        product_sheet.append(["Date/time", "Product", "Unique users", "Monthly cost (EUR)"])
    elif product_sheet.max_row == 1 and product_sheet.cell(1, 1).value != "Date/time":
        product_sheet.delete_rows(1, product_sheet.max_row)
        product_sheet.append(["Date/time", "Product", "Unique users", "Monthly cost (EUR)"])
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for row in summary:
        product_sheet.append([timestamp, row["product"], row["unique_users"], row["cost"] or 0])
    for cell in product_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E20074")
    for row in product_sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = '€#,##0.00'
    for column_cells in product_sheet.columns:
        width = min(max(len(clean(c.value)) for c in column_cells) + 2, 42)
        product_sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    product_sheet.freeze_panes = "A2"
    wb.save(history_path)


def add_history_charts(workbook):
    """Add a compact overview and daily trends to an exported history workbook."""
    for sheet_name in ("Overview", "Trends", "Charts & Trends"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
    history = workbook["History"]
    if history.max_row < 2:
        return

    records = []
    for row in history.iter_rows(min_row=2, values_only=True):
        values = list(row) + [None] * 8
        records.append(values[:8])
    latest_by_day = {}
    for row in records:
        timestamp = str(row[0] or "")
        latest_by_day[timestamp[:10]] = row
    daily_rows = [latest_by_day[key] for key in sorted(latest_by_day)]

    overview = workbook.create_sheet("Overview", 0)
    overview["A1"] = "T-Digital AI Cost Calculator"
    overview["A2"] = "History overview"
    overview.append([])
    overview.append(["Metric", "Latest value"])
    latest = daily_rows[-1]
    metrics = [
        ("Latest execution", latest[0]),
        ("Source report", latest[1]),
        ("Unique AI users", latest[2]),
        ("Monthly cost (EUR)", latest[3]),
        ("Average cost/user (EUR)", latest[5]),
        ("12-month projection (EUR)", latest[4]),
        ("Users without AI license", latest[7]),
        ("Tracked days", len(daily_rows)),
    ]
    for metric, value in metrics:
        overview.append([metric, value])
    for cell in overview[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E20074")
    overview.column_dimensions["A"].width = 32
    overview.column_dimensions["B"].width = 28
    for row_number in (8, 9, 10):
        overview.cell(row_number, 2).number_format = '€#,##0.00'
    overview.freeze_panes = "A5"

    trends = workbook.create_sheet("Trends", 1)
    trends["A1"] = "Daily trends"
    trends["A2"] = "One point per day: the last saved execution of that day."
    trends.append([])
    trends.append(["Date", "Monthly cost (EUR)", "Unique AI users", "Users without AI license"])
    for row in daily_rows:
        trends.append([row[0], row[3], row[2], row[7]])
    for cell in trends[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E20074")
    for column, width in {"A": 20, "B": 22, "C": 18, "D": 25}.items():
        trends.column_dimensions[column].width = width
    trends.freeze_panes = "A5"

    def line_chart(title, min_col, max_col, position, y_title):
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Date"
        chart.height = 7
        chart.width = 16
        data = Reference(trends, min_col=min_col, max_col=max_col, min_row=4, max_row=trends.max_row)
        categories = Reference(trends, min_col=1, min_row=5, max_row=trends.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        trends.add_chart(chart, position)

    line_chart("Monthly cost", 2, 2, "F4", "EUR")
    line_chart("AI users and users without license", 3, 4, "F20", "Users")


def show_history(root):
    window = tk.Toplevel(root)
    window.title("AI Cost Calculator - History")
    window.geometry("1160x560")
    window.minsize(800, 400)
    window.transient(root)
    window.columnconfigure(0, weight=1)
    window.rowconfigure(1, weight=1)
    history_path = history_file()
    if not history_path.exists():
        ttk.Label(window, text="No history available yet.", padding=30).grid(
            row=0, column=0, sticky="nsew"
        )
        return

    workbook = load_workbook(history_path, data_only=True, read_only=True)
    notebook = ttk.Notebook(window)
    notebook.grid(row=0, column=0, rowspan=2, sticky="nsew", padx=12, pady=(10, 0))

    def build_history_tab(sheet_name, title, headings):
        tab = ttk.Frame(notebook, padding=8)
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(1, weight=1)
        notebook.add(tab, text=title)
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else None
        rows = list(sheet.iter_rows(values_only=True)) if sheet else []
        ttk.Label(tab, text=f"{max(0, len(rows) - 1)} saved calculations").grid(
            row=0, column=0, sticky="w", pady=(0, 6)
        )
        if not rows:
            ttk.Label(tab, text="No history available yet.", padding=30).grid(
                row=1, column=0, sticky="nsew"
            )
            return
        columns = [f"column_{index}" for index in range(len(headings))]
        table_frame = ttk.Frame(tab)
        table_frame.grid(row=1, column=0, sticky="nsew")
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
        widths = [155, 260, 135, 150, 175, 150, 205, 175]
        for index, (key, heading) in enumerate(zip(columns, headings)):
            tree.heading(key, text=heading)
            anchor = "center" if index in {2, 3, 4, 5, 7} else "w"
            tree.column(key, width=widths[index], minwidth=90, anchor=anchor, stretch=False)
        for row in rows[1:]:
            values = list(row) + [""] * len(headings)
            tree.insert("", "end", values=values[:len(headings)])
        vertical_scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        horizontal_scrollbar = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vertical_scrollbar.set, xscrollcommand=horizontal_scrollbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vertical_scrollbar.grid(row=0, column=1, sticky="ns")
        horizontal_scrollbar.grid(row=1, column=0, sticky="ew")

    headings = ["Date/time", "Source report", "Unique users", "Monthly cost (EUR)", "12-month projection (EUR)", "Average/user (EUR)", "Costcenter data updated", "Users without license"]
    build_history_tab("History", "AI History", headings)
    build_history_tab("Full License History", "Full License History", headings)

    def export_history():
        destination = filedialog.asksaveasfilename(
            parent=window,
            title="Export history",
            initialdir=str(export_directory()),
            initialfile=f"AI_Cost_History_{datetime.now():%Y%m%d_%H%M}.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if destination:
            shutil.copy2(history_path, destination)
            exported = load_workbook(destination)
            add_history_charts(exported)
            exported.save(destination)
            messagebox.showinfo("History exported", f"Saved to:\n{destination}", parent=window)

    button_bar = ttk.Frame(window, padding=(12, 8, 12, 12))
    button_bar.grid(row=2, column=0, sticky="ew")
    button_bar.columnconfigure(0, weight=1)
    ttk.Button(button_bar, text="Export History", command=export_history).grid(
        row=0, column=1, padx=(6, 0)
    )
    ttk.Button(button_bar, text="Close", command=window.destroy).grid(
        row=0, column=2, padx=(6, 0)
    )


def choose_costcenter(root, members):
    costcenters = sorted({m.get("costcenter", "") for m in members.values() if m.get("costcenter")})
    if not costcenters:
        messagebox.showerror("No cost centers", "No cost centers were found in the saved member report.", parent=root)
        return None
    window = tk.Toplevel(root)
    window.title("Select Cost Center")
    window.geometry("420x150")
    window.transient(root)
    window.grab_set()
    selected = tk.StringVar(value=costcenters[0])
    ttk.Label(window, text="Cost center:").pack(pady=(20, 6))
    combo = ttk.Combobox(window, textvariable=selected, values=costcenters, state="readonly", width=35)
    combo.pack()
    result = {"value": None}

    def confirm():
        result["value"] = selected.get()
        window.destroy()

    ttk.Button(window, text="Continue", command=confirm).pack(pady=14)
    root.wait_window(window)
    return result["value"]


def choose_manager(root, members):
    managers = sorted({m.get("manager", "") for m in members.values() if m.get("manager")}, key=str.casefold)
    if not managers:
        messagebox.showerror(
            "No managers",
            "No cost center managers were found in the saved member report.",
            parent=root,
        )
        return None
    window = tk.Toplevel(root)
    window.title("Select Cost Center Manager")
    window.geometry("520x150")
    window.transient(root)
    window.grab_set()
    selected = tk.StringVar(value=managers[0])
    ttk.Label(window, text="Cost center manager:").pack(pady=(20, 6))
    combo = ttk.Combobox(window, textvariable=selected, values=managers, state="readonly", width=45)
    combo.pack()
    result = {"value": None}

    def confirm():
        result["value"] = selected.get()
        window.destroy()

    ttk.Button(window, text="Continue", command=confirm).pack(pady=14)
    root.wait_window(window)
    return result["value"]


def choose_licenses(root, services):
    available = sorted({item["service"] for item in services if item["service"] in active_license_names()})
    if not available:
        messagebox.showerror("No AI licenses", "No configured AI licenses were found in the report.", parent=root)
        return None
    window = tk.Toplevel(root)
    window.title("Select Licenses")
    window.geometry("520x360")
    window.transient(root)
    window.grab_set()
    ttk.Label(window, text="Select one or more licenses:").pack(pady=(16, 8))
    listbox = tk.Listbox(window, selectmode=tk.MULTIPLE, height=12, width=62)
    for license_name in available:
        listbox.insert(tk.END, license_name)
    listbox.pack(padx=16, fill="both", expand=True)
    result = {"value": None}

    def confirm():
        result["value"] = [available[index] for index in listbox.curselection()]
        window.destroy()

    ttk.Button(window, text="Continue", command=confirm).pack(pady=12)
    root.wait_window(window)
    return result["value"]


def ask_for_emails(root):
    window = tk.Toplevel(root)
    window.title("Enter User Emails")
    window.geometry("540x350")
    window.transient(root)
    window.grab_set()
    ttk.Label(window, text="Paste emails, one per line or separated by commas:").pack(pady=(16, 8))
    text_area = tk.Text(window, height=12, width=62)
    text_area.pack(padx=16, fill="both", expand=True)
    result = {"emails": None}

    def confirm():
        raw = text_area.get("1.0", "end")
        values = raw.replace(",", "\n").splitlines()
        result["emails"] = sorted({value.strip().casefold() for value in values if value.strip()})
        window.destroy()

    ttk.Button(window, text="Continue", command=confirm).pack(pady=12)
    root.wait_window(window)
    return result["emails"]


def manage_licenses(root):
    window = tk.Toplevel(root)
    window.title("License Settings")
    window.geometry("780x500")
    window.minsize(650, 400)
    window.transient(root)
    window.grab_set()
    ttk.Label(
        window,
        text="Manage AI and non-AI licenses used by future calculations.",
    ).pack(anchor="w", padx=16, pady=(16, 8))

    notebook = ttk.Notebook(window)
    notebook.pack(fill="both", expand=True, padx=16, pady=4)

    def build_tab(title, catalog, active_map):
        tab = ttk.Frame(notebook, padding=8)
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(0, weight=1)
        notebook.add(tab, text=title)
        columns = ("license", "price", "status")
        table = ttk.Treeview(tab, columns=columns, show="headings", height=13)
        table.heading("license", text="License")
        table.heading("price", text="Price / user / month (€)")
        table.heading("status", text="Status")
        table.column("license", width=420)
        table.column("price", width=170, anchor="e")
        table.column("status", width=110, anchor="center")
        scrollbar = ttk.Scrollbar(tab, orient="vertical", command=table.yview)
        table.configure(yscrollcommand=scrollbar.set)
        table.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        def refresh():
            for item in table.get_children():
                table.delete(item)
            for name in sorted(catalog, key=str.casefold):
                _, price = catalog[name]
                status = "Active" if active_map.get(name, True) else "Inactive"
                table.insert("", tk.END, iid=name, values=(name, f"€{price:,.2f}", status))

        def selected_name():
            selected = table.selection()
            return selected[0] if selected else None

        def add_license():
            name = simpledialog.askstring("Add license", "License name:", parent=window)
            if not name:
                return
            name = name.strip()
            if name in PRODUCT_PRICES or name in NON_AI_PRODUCT_PRICES:
                messagebox.showerror("License exists", "A license with this name already exists.", parent=window)
                return
            price = simpledialog.askfloat("Add license", "Monthly price per user (€):", parent=window, minvalue=0)
            if price is None:
                return
            catalog[name] = (name, price)
            active_map[name] = True
            save_license_settings()
            refresh()

        def edit_license():
            name = selected_name()
            if not name:
                messagebox.showinfo("Select license", "Select a license first.", parent=window)
                return
            _, old_price = catalog[name]
            price = simpledialog.askfloat(
                "Change license price", f"Monthly price for {name} (€):",
                initialvalue=old_price, minvalue=0, parent=window,
            )
            if price is not None:
                product, _ = catalog[name]
                catalog[name] = (product, price)
                save_license_settings()
                refresh()

        def toggle_license():
            name = selected_name()
            if not name:
                messagebox.showinfo("Select license", "Select a license first.", parent=window)
                return
            active = active_map.get(name, True)
            action = "deactivate" if active else "reactivate"
            if not messagebox.askyesno(
                f"{action.title()} license",
                f"{action.title()} '{name}' for future calculations?\n\n"
                "Old exports and history will not be changed.", parent=window,
            ):
                return
            active_map[name] = not active
            save_license_settings()
            refresh()

        buttons = ttk.Frame(tab)
        buttons.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Button(buttons, text="Add License", command=add_license).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Change Price", command=edit_license).pack(side="left", padx=6)
        ttk.Button(buttons, text="Deactivate / Reactivate", command=toggle_license).pack(side="left", padx=6)
        refresh()

    build_tab("AI Licenses", PRODUCT_PRICES, LICENSE_ACTIVE)
    build_tab("Non-AI Licenses", NON_AI_PRODUCT_PRICES, NON_AI_LICENSE_ACTIVE)
    ttk.Button(window, text="Close", command=window.destroy).pack(anchor="e", padx=16, pady=(6, 14))
    root.wait_window(window)


def add_daily_joke_ticker(parent):
    setups = [
        "Why did the spreadsheet apply for a job?",
        "Why was the dashboard invited to the meeting?",
        "Why did the analyst bring a ladder to work?",
        "Why did the report go to therapy?",
        "Why did the database take a coffee break?",
        "Why did the formula get promoted?",
        "Why did the pivot table feel confident?",
        "Why did the project plan stay calm?",
        "Why did the keyboard join the team?",
        "Why did the chart get applause?",
        "Why did the data scientist bring an umbrella?",
        "Why did the server go to the gym?",
        "Why did the calendar become a manager?",
        "Why did the analyst check the numbers twice?",
        "Why did the laptop ask for a holiday?",
        "Why did the meeting room need a map?",
        "Why did the spreadsheet wear a tie?",
        "Why did the cloud get a compliment?",
        "Why did the query cross the office?",
        "Why did the robot bring a notebook?",
        "Why did the team invite the calculator?",
        "Why did the chart avoid gossip?",
        "Why did the data pipeline take the stairs?",
        "Why did the inbox start exercising?",
        "Why did the analyst bring a ruler?",
        "Why did the report arrive early?",
        "Why did the meeting agenda smile?",
        "Why did the spreadsheet open a savings account?",
        "Why did the code review bring snacks?",
        "Why did the dashboard wear sunglasses?",
        "Why did the metric get a certificate?",
        "Why did the file name stay short?",
        "Why did the server bring a jacket?",
        "Why did the team use a bookmark?",
        "Why did the chart bring a microphone?",
        "Why did the formula take notes?",
        "Why did the laptop join the coffee queue?",
        "Why did the database tell a story?",
        "Why did the analyst bring a compass?",
        "Why did the project plan bring a stopwatch?",
        "Why did the spreadsheet get a thank-you note?",
        "Why did the query ask for directions?",
        "Why did the calendar book a meeting with itself?",
        "Why did the report carry an umbrella?",
        "Why did the dashboard bring a ruler?",
        "Why did the team give the server a high five?",
        "Why did the code bring a lunchbox?",
        "Why did the number wear a badge?",
        "Why did the chart bring a telescope?",
        "Why did the AI assistant stay relaxed?",
    ]
    punchlines = [
        "It wanted to improve its cell-f.",
        "It was looking for better data and fewer dramas.",
        "Because every good answer starts with a solid calculation.",
        "It had excellent processing under pressure.",
        "It knew the best solution was just one refresh away.",
        "It wanted to reach the next level of insights.",
        "It believed teamwork makes the workflow work.",
        "It preferred clear outputs and short meetings.",
        "It was ready to turn a problem into a useful report.",
        "It knew that even small improvements add up.",
    ]
    jokes = [f"{setup} {punchline}" for setup in setups for punchline in punchlines]
    random.SystemRandom().shuffle(jokes)
    joke_index = 0
    ticker = tk.Frame(parent, background="#ffffff", height=32)
    ticker.pack(fill="x", padx=34, pady=(0, 10))
    ticker.pack_propagate(False)
    canvas = tk.Canvas(ticker, background="#ffffff", highlightthickness=0, height=32)
    canvas.pack(fill="both", expand=True)
    canvas.create_text(10, 16, text="DAILY NOTE", fill="#e20074", anchor="w", font=("Arial", 9, "bold"))
    text_id = canvas.create_text(105, 16, text=jokes[joke_index], fill="#555555", anchor="w", font=("Arial", 10))

    def new_joke():
        nonlocal joke_index
        joke_index = (joke_index + 1) % len(jokes)
        canvas.itemconfigure(text_id, text=jokes[joke_index])
        canvas.coords(text_id, 105, 16)

    ttk.Button(parent, text="New joke", command=new_joke).pack(anchor="e", padx=34, pady=(0, 8))

    def scroll():
        canvas.move(text_id, -1, 0)
        bounds = canvas.bbox(text_id)
        if bounds and bounds[2] < 0:
            canvas.move(text_id, canvas.winfo_width() - bounds[0] + 20, 0)
        canvas.after(35, scroll)

    canvas.after(300, scroll)


def show_startup_splash(root):
    """Show a short T-Digital splash screen only while the app starts."""
    root.withdraw()
    splash = tk.Toplevel(root)
    splash.overrideredirect(True)
    splash.configure(background="#e20074")
    splash.geometry("420x220")
    splash.update_idletasks()
    screen_width = splash.winfo_screenwidth()
    screen_height = splash.winfo_screenheight()
    x = (screen_width - splash.winfo_width()) // 2
    y = (screen_height - splash.winfo_height()) // 2
    splash.geometry(f"+{x}+{y}")
    tk.Label(
        splash,
        text="T",
        background="#e20074",
        foreground="white",
        font=("Arial", 54, "bold"),
    ).pack(pady=(24, 0))
    tk.Label(
        splash,
        text="T-Digital AI Cost Calculator",
        background="#e20074",
        foreground="white",
        font=("Arial", 17, "bold"),
    ).pack(pady=(4, 0))
    tk.Label(
        splash,
        text="License cost management",
        background="#e20074",
        foreground="white",
        font=("Arial", 10),
    ).pack(pady=(5, 0))

    def finish():
        if splash.winfo_exists():
            splash.destroy()
        root.deiconify()
        root.lift()

    splash.after(1500, finish)


def configure_windows_dpi():
    """Let Windows provide logical pixels so the UI is correct at any DPI."""
    if sys.platform != "win32":
        return
    try:
        import ctypes

        ctypes.windll.user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))
    except (AttributeError, OSError):
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except (AttributeError, OSError):
            pass


def main():
    load_license_settings()
    configure_windows_dpi()
    root = tk.Tk()
    root.title("T-Digital AI Cost Calculator")
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    initial_width = min(max(int(screen_width * 0.38), 560), 760)
    initial_height = min(max(int(screen_height * 0.72), 620), 900)
    root.geometry(f"{initial_width}x{initial_height}")
    root.minsize(min(560, max(420, screen_width - 80)), min(620, max(460, screen_height - 100)))
    root.resizable(True, True)
    root.configure(background="#f3f3f3")
    show_startup_splash(root)
    style = ttk.Style(root)
    style.configure("Telekom.TButton", font=("Arial", 11), padding=(12, 8))
    style.configure("Telekom.TLabel", background="#f3f3f3", foreground="#333333")
    main_canvas = tk.Canvas(root, background="#f3f3f3", highlightthickness=0)
    scrollbar = ttk.Scrollbar(root, orient="vertical", command=main_canvas.yview)
    main_canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")
    main_canvas.configure(yscrollcommand=scrollbar.set)
    content = ttk.Frame(main_canvas)
    content_id = main_canvas.create_window((0, 0), window=content, anchor="nw")
    content.bind("<Configure>", lambda event: main_canvas.configure(scrollregion=main_canvas.bbox("all")))
    main_canvas.bind("<Configure>", lambda event: main_canvas.itemconfigure(content_id, width=event.width))
    main_canvas.bind_all("<MouseWheel>", lambda event: main_canvas.yview_scroll(int(-event.delta / 120), "units"))
    main_canvas.bind_all("<Button-4>", lambda event: main_canvas.yview_scroll(-3, "units"))
    main_canvas.bind_all("<Button-5>", lambda event: main_canvas.yview_scroll(3, "units"))
    header = tk.Frame(content, background="#e20074", height=108)
    header.pack(fill="x")
    header.pack_propagate(False)
    tk.Label(header, text="T", background="#e20074", foreground="white", font=("Arial", 34, "bold")).pack(side="left", padx=(26, 12))
    title_frame = tk.Frame(header, background="#e20074")
    title_frame.pack(side="left", pady=18)
    tk.Label(title_frame, text="T-Digital AI Cost Calculator", background="#e20074", foreground="white", font=("Arial", 21, "bold"), anchor="w").pack(anchor="w")
    tk.Label(title_frame, text="License cost management", background="#e20074", foreground="white", font=("Arial", 11), anchor="w").pack(anchor="w")
    add_daily_joke_ticker(content)
    ttk.Label(content, text="Calculate monthly AI licence costs and keep a history.", style="Telekom.TLabel").pack(pady=(18, 10))
    current_members, member_updated = load_saved_members()
    status_var = tk.StringVar(value=f"Costcenter data: {member_updated}")
    ttk.Label(content, textvariable=status_var, foreground="#555555").pack(pady=(0, 12))
    current_source = None
    source_status_var = tk.StringVar(value="AD report: Not selected")
    ttk.Label(content, textvariable=source_status_var, foreground="#555555").pack(pady=(0, 12))
    preview_vars = {
        "users": tk.StringVar(value="—"),
        "monthly": tk.StringVar(value="—"),
        "average": tk.StringVar(value="—"),
        "projection": tk.StringVar(value="—"),
        "without": tk.StringVar(value="—"),
    }
    preview = ttk.LabelFrame(content, text="Live Preview", padding=10)
    preview.pack(fill="x", padx=34, pady=(0, 14))
    preview_columns = [
        ("users", "Unique AI users"),
        ("monthly", "Monthly cost"),
        ("average", "Cost / user"),
        ("projection", "12-month projection"),
        ("without", "Users without license"),
    ]
    for column, (key, label) in enumerate(preview_columns):
        card = ttk.Frame(preview)
        card.grid(row=0, column=column, sticky="nsew", padx=4)
        preview.columnconfigure(column, weight=1)
        ttk.Label(card, text=label, style="Telekom.TLabel", wraplength=120, justify="center").pack()
        ttk.Label(card, textvariable=preview_vars[key], font=("Arial", 13, "bold"), foreground="#e20074").pack(pady=(4, 0))

    def update_preview(summary, unique_users, without_license):
        total_cost = sum(row["cost"] or 0 for row in summary)
        user_count = len(unique_users)
        preview_vars["users"].set(f"{user_count:,}")
        preview_vars["monthly"].set(f"€{total_cost:,.2f}")
        preview_vars["average"].set(f"€{total_cost / user_count:,.2f}" if user_count else "n.a.")
        preview_vars["projection"].set(f"€{total_cost * 12:,.2f}")
        preview_vars["without"].set(f"{len(without_license):,}")

    def update_members():
        nonlocal current_members, member_updated
        member_name = filedialog.askopenfilename(
            parent=root,
            title="Select Costcenter Member Report",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if not member_name:
            return False
        try:
            current_members = read_costcenter_report(Path(member_name))
            member_updated = save_members(current_members)
            status_var.set(f"Costcenter data: {member_updated}")
            messagebox.showinfo(
                "Costcenter data updated",
                f"Saved {len(current_members)} members.\n\n"
                "This report will be used automatically for future calculations.",
                parent=root,
            )
            return True
        except Exception as exc:
            messagebox.showerror("Could not update Costcenter data", str(exc), parent=root)
            return False

    def select_source_report():
        nonlocal current_source
        source_name = filedialog.askopenfilename(
            parent=root,
            title="Select AD Based License Report",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if not source_name:
            return None
        current_source = Path(source_name)
        source_status_var.set(f"AD report: {current_source.name}")
        return current_source

    def get_source_report():
        if current_source and current_source.exists():
            return current_source
        return select_source_report()

    def run_calculator():
        source = get_source_report()
        if not source:
            return
        try:
            if not current_members and not update_members():
                return
            services = read_report(source)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(services, current_members)
            if not detail:
                raise ValueError("No matching AI services were found.")
            update_preview(summary, unique_users, without_license)
            default_name = f"AI_License_Cost_Summary_{datetime.now():%Y%m%d_%H%M}.xlsx"
            destination_dir = export_directory()
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save cost export",
                initialdir=str(destination_dir),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail, unique_users,
                    user_costs, without_license,
                    user_licenses,
                )
                append_history(source, summary, unique_users, len(without_license), member_updated)
                total_cost = sum(row["cost"] or 0 for row in summary)
                average = total_cost / len(unique_users) if unique_users else 0
                messagebox.showinfo(
                    "Export complete",
                    f"Saved to:\n{destination_name}\n\n"
                    f"Unique AI users: {len(unique_users)}\n"
                    f"Monthly cost: €{total_cost:,.2f}\n"
                    f"12-month projection: €{total_cost * 12:,.2f}\n"
                    f"Average/user: €{average:,.2f}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not process report", str(exc), parent=root)

    def run_full_license_export():
        source = get_source_report()
        if not source:
            return
        try:
            if not current_members and not update_members():
                return
            services = read_report(source)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                services, current_members, include_non_ai=True
            )
            if not detail:
                raise ValueError("No configured licenses were found in the report.")
            default_name = f"Full_License_Report_{datetime.now():%Y%m%d_%H%M}.xlsx"
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save full license report",
                initialdir=str(export_directory()),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail, unique_users,
                    user_costs, without_license, user_licenses,
                    full_license_report=True,
                )
                append_history(
                    source, summary, unique_users, len(without_license), member_updated,
                    history_sheet_name="Full License History",
                    product_sheet_name="Full Product History",
                    full_license_report=True,
                )
                messagebox.showinfo(
                    "Full license export complete",
                    f"Saved to:\n{destination_name}\n\n"
                    f"Unique licensed users: {len(unique_users)}\n"
                    f"Monthly cost: €{sum(row['cost'] or 0 for row in summary):,.2f}\n"
                    f"Users without any license: {len(without_license)}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export full license report", str(exc), parent=root)

    def run_costcenter_export():
        nonlocal current_members, member_updated
        if not current_members and not update_members():
            return
        source = get_source_report()
        if not source:
            return
        costcenter = choose_costcenter(root, current_members)
        if not costcenter:
            return
        try:
            services = read_report(source)
            selected_members = {
                key: value for key, value in current_members.items()
                if value.get("costcenter") == costcenter
            }
            selected_accounts = set(selected_members)
            selected_services = []
            for item in services:
                users = [u for u in item["users"] if (u["account"] or u["username"]).casefold() in selected_accounts]
                if users:
                    selected_item = dict(item)
                    selected_item["users"] = users
                    selected_item["reported_total"] = len(users)
                    selected_services.append(selected_item)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                selected_services, selected_members
            )
            if not detail:
                raise ValueError(f"No AI licences found for cost center {costcenter}.")
            default_name = f"AI_License_Cost_{costcenter}_{datetime.now():%Y%m%d_%H%M}.xlsx"
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save cost center export",
                initialdir=str(export_directory()),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(Path(destination_name), source, summary, detail, unique_users, user_costs, without_license, user_licenses)
                messagebox.showinfo(
                    "Cost center export complete",
                    f"Cost center: {costcenter}\n"
                    f"Unique AI users: {len(unique_users)}\n"
                    f"Monthly cost: €{sum(row['cost'] or 0 for row in summary):,.2f}\n"
                    f"Saved to:\n{destination_name}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export cost center", str(exc), parent=root)

    def run_manager_export():
        nonlocal current_members, member_updated
        if not current_members and not update_members():
            return
        source = get_source_report()
        if not source:
            return
        manager = choose_manager(root, current_members)
        if not manager:
            return
        try:
            services = read_report(source)
            selected_members = {
                key: value for key, value in current_members.items()
                if value.get("manager") == manager
            }
            selected_accounts = set(selected_members)
            selected_services = []
            for item in services:
                users = [
                    u for u in item["users"]
                    if (u["account"] or u["username"]).casefold() in selected_accounts
                ]
                if users:
                    selected_item = dict(item)
                    selected_item["users"] = users
                    selected_item["reported_total"] = len(users)
                    selected_services.append(selected_item)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                selected_services, selected_members
            )
            if not detail:
                raise ValueError(f"No AI licences found for manager {manager}.")
            safe_manager = "_".join(manager.split())[:60]
            default_name = f"AI_License_Cost_Manager_{safe_manager}_{datetime.now():%Y%m%d_%H%M}.xlsx"
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save manager export",
                initialdir=str(export_directory()),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail,
                    unique_users, user_costs, without_license, user_licenses,
                )
                messagebox.showinfo(
                    "Manager export complete",
                    f"Manager: {manager}\n"
                    f"Unique AI users: {len(unique_users)}\n"
                    f"Monthly cost: €{sum(row['cost'] or 0 for row in summary):,.2f}\n"
                    f"Saved to:\n{destination_name}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export manager", str(exc), parent=root)

    def run_email_export():
        nonlocal current_members, member_updated
        if not current_members and not update_members():
            return
        requested_emails = ask_for_emails(root)
        if not requested_emails:
            return
        source = get_source_report()
        if not source:
            return
        try:
            email_to_account = {
                value.get("email", "").casefold(): key
                for key, value in current_members.items()
                if value.get("email")
            }
            selected_accounts = {email_to_account[email] for email in requested_emails if email in email_to_account}
            missing_emails = [email for email in requested_emails if email not in email_to_account]
            selected_members = {key: current_members[key] for key in selected_accounts}
            services = read_report(source)
            selected_services = []
            for item in services:
                users = [u for u in item["users"] if (u["account"] or u["username"]).casefold() in selected_accounts]
                if users:
                    selected_item = dict(item)
                    selected_item["users"] = users
                    selected_item["reported_total"] = len(users)
                    selected_services.append(selected_item)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                selected_services, selected_members
            )
            default_name = f"AI_License_Cost_Selected_Users_{datetime.now():%Y%m%d_%H%M}.xlsx"
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save selected users export",
                initialdir=str(export_directory()),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail,
                    unique_users, user_costs, without_license, user_licenses,
                    missing_emails,
                )
                messagebox.showinfo(
                    "Selected users export complete",
                    f"Emails requested: {len(requested_emails)}\n"
                    f"Users found: {len(selected_members)}\n"
                    f"Users with AI license: {len(unique_users)}\n"
                    f"Emails not found: {len(missing_emails)}\n\n"
                    f"Saved to:\n{destination_name}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export selected users", str(exc), parent=root)

    def run_multiple_license_export():
        nonlocal current_members, member_updated
        if not current_members and not update_members():
            return
        source = get_source_report()
        if not source:
            return
        try:
            services = read_report(source)
            _, _, _, all_user_costs, _, _ = calculate(services, current_members)
            selected_accounts = {
                user["account"].casefold() for user in all_user_costs
                if user["chargeable_licenses"] > 1
            }
            if not selected_accounts:
                messagebox.showinfo("No matching users", "No users have more than one chargeable AI license.", parent=root)
                return
            selected_members = {key: current_members[key] for key in selected_accounts if key in current_members}
            selected_services = []
            for item in services:
                users = [u for u in item["users"] if (u["account"] or u["username"]).casefold() in selected_accounts]
                if users:
                    selected_item = dict(item)
                    selected_item["users"] = users
                    selected_item["reported_total"] = len(users)
                    selected_services.append(selected_item)
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                selected_services, selected_members
            )
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save multiple-license users export",
                initialdir=str(export_directory()),
                initialfile=f"AI_Users_Multiple_Chargeable_Licenses_{datetime.now():%Y%m%d_%H%M}.xlsx",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail, unique_users,
                    user_costs, without_license, user_licenses,
                )
                messagebox.showinfo(
                    "Export complete",
                    f"Users with more than one chargeable license: {len(unique_users)}\n\n"
                    f"Saved to:\n{destination_name}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export multiple-license users", str(exc), parent=root)

    def run_license_export():
        nonlocal current_members, member_updated
        if not current_members and not update_members():
            return
        source = get_source_report()
        if not source:
            return
        try:
            services = read_report(source)
            selected_names = choose_licenses(root, services)
            if not selected_names:
                return
            selected_services = [item for item in services if item["service"] in selected_names]
            summary, detail, unique_users, user_costs, without_license, user_licenses = calculate(
                selected_services, current_members
            )
            destination_name = filedialog.asksaveasfilename(
                parent=root,
                title="Save selected licenses export",
                initialdir=str(export_directory()),
                initialfile=f"AI_License_Cost_Selected_Licenses_{datetime.now():%Y%m%d_%H%M}.xlsx",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
            )
            if destination_name:
                export_report(
                    Path(destination_name), source, summary, detail,
                    unique_users, user_costs, without_license, user_licenses,
                )
                messagebox.showinfo(
                    "Selected licenses export complete",
                    f"Licenses selected: {len(selected_names)}\n"
                    f"Unique AI users: {len(unique_users)}\n"
                    f"Saved to:\n{destination_name}",
                    parent=root,
                )
        except Exception as exc:
            messagebox.showerror("Could not export selected licenses", str(exc), parent=root)

    ttk.Label(content, text="EXPORTS", style="Telekom.TLabel", font=("Arial", 10, "bold")).pack(anchor="w", padx=34, pady=(6, 2))
    exports = ttk.Frame(content)
    exports.pack(fill="x", padx=34)
    exports.columnconfigure(0, weight=1)
    exports.columnconfigure(1, weight=1)
    ttk.Button(exports, text="Calculate from Excel report", style="Telekom.TButton", command=run_calculator).grid(row=0, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export Cost Center", style="Telekom.TButton", command=run_costcenter_export).grid(row=1, column=0, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export by User Emails", style="Telekom.TButton", command=run_email_export).grid(row=1, column=1, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export by Cost Center Manager", style="Telekom.TButton", command=run_manager_export).grid(row=2, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export by Licenses", style="Telekom.TButton", command=run_license_export).grid(row=3, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export Users with Multiple Chargeable Licenses", style="Telekom.TButton", command=run_multiple_license_export).grid(row=4, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(exports, text="Export Full License Report", style="Telekom.TButton", command=run_full_license_export).grid(row=5, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Label(content, text="DATA & HISTORY", style="Telekom.TLabel", font=("Arial", 10, "bold")).pack(anchor="w", padx=34, pady=(12, 2))
    data_frame = ttk.Frame(content)
    data_frame.pack(fill="x", padx=34)
    data_frame.columnconfigure(0, weight=1)
    data_frame.columnconfigure(1, weight=1)
    ttk.Button(data_frame, text="Select / Change AD Report", style="Telekom.TButton", command=select_source_report).grid(row=0, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(data_frame, text="Update Costcenter Members", style="Telekom.TButton", command=update_members).grid(row=1, column=0, sticky="ew", padx=3, pady=3)
    ttk.Button(data_frame, text="View History", style="Telekom.TButton", command=lambda: show_history(root)).grid(row=1, column=1, sticky="ew", padx=3, pady=3)
    ttk.Button(data_frame, text="License Settings", style="Telekom.TButton", command=lambda: manage_licenses(root)).grid(row=2, column=0, columnspan=2, sticky="ew", padx=3, pady=3)
    ttk.Button(content, text="Exit", style="Telekom.TButton", command=root.destroy).pack(fill="x", padx=37, pady=(16, 6))
    root.mainloop()


if __name__ == "__main__":
    main()
